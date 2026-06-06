from openpyxl import Workbook
from openpyxl.styles import Font
from database import get_user_expenses

def export_to_excel(user_id):
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"

    headers = [
        "ID",
        "مبلغ",
        "دسته‌بندی",
        "محل",
        "تاریخ شمسی",
        "تاریخ میلادی",
        "توضیحات"
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = Font(bold=True)

    expenses = get_user_expenses(user_id)

    for row_num, expense in enumerate(expenses, 2):
        ws.cell(row=row_num, column=1).value = expense[0]
        ws.cell(row=row_num, column=2).value = expense[2]
        ws.cell(row=row_num, column=3).value = expense[3]
        ws.cell(row=row_num, column=4).value = expense[4]
        ws.cell(row=row_num, column=5).value = expense[5]
        ws.cell(row=row_num, column=6).value = expense[6]
        ws.cell(row=row_num, column=7).value = expense[7]

    filename = f"expenses_{user_id}.xlsx"
    wb.save(filename)

    return filename